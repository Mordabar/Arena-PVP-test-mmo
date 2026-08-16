#!/usr/bin/env python3
"""Independent adversarial audit for v0.13 source-power parity.

Does NOT import the generator. It re-reads the DOCX independently and compares
its source contract with the generated runtime catalogue field by field.
"""
from __future__ import annotations
from pathlib import Path
from docx import Document
import argparse, json, re, sys, unicodedata

TARGETS={
 'GUERRERO':['devastador','guardian'],'CABALLERO':['guardian'],'BÁRBARO':['devastador'],
 'ARQUERO':['centinela','rastreador'],'CAZADOR':['rastreador'],'TIRADOR':['centinela'],
 'MAGO':['arcanista','vinculador'],'BRUJO':['arcanista'],'CONJURADOR':['vinculador']}
SUFFIX={
 'GUERRERO':'de Guerra','CABALLERO':'del Bastión','BÁRBARO':'de Sangre','ARQUERO':'del Viento',
 'CAZADOR':'del Acecho','TIRADOR':'del Horizonte','MAGO':'del Éter','BRUJO':'del Vacío','CONJURADOR':'del Nexo'}
GCD={'Muy corto':'reactive','Corto':'short','Normal':'standard','Largo':0.8,'Muy largo':0.9,'-':'none'}
NUM=re.compile(r'-?\d+(?:[.,]\d+)?')

def strip_paren(s): return re.sub(r'\s*\([^)]*\)\s*$','',s or '').strip()
def lastnum(s, default=0.0):
    xs=NUM.findall(s or ''); return float(xs[-1].replace(',','.')) if xs else float(default)
def rank5(s):
    v=(s or '').split(':',1)[1].strip() if ':' in (s or '') else (s or '').strip()
    p=re.split(r'\s+/\s+',v); return (p[-1] if p else v).strip()
def boolish(s):
    l=(s or '').lower(); return ('sí' in l or 'yes' in l or '100%' in l)
def norm(s):
    s=unicodedata.normalize('NFKD',s or '').encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+',' ',s).strip()

def parse_source(path:Path):
    doc=Document(str(path)); heads=[]; cls=None; disc=None; ps=doc.paragraphs
    valid_classes=set(TARGETS)
    for i,p in enumerate(ps):
        st=p.style.name if p.style else ''; t=p.text.strip()
        if st=='Heading 1' and t:
            k=strip_paren(t).upper()
            if k in valid_classes: cls=k; disc=None
        elif st=='Heading 2' and t and cls: disc=strip_paren(t)
        elif st=='Heading 3' and cls:
            bullets=[]; section=''; desc=''; j=i+1
            while j<len(ps):
                q=ps[j]; qs=q.style.name if q.style else ''; qt=q.text.strip()
                if qs in ('Heading 1','Heading 2','Heading 3'): break
                if qt:
                    if qt in ('Daño','Beneficios (buffs)','Efectos negativos / Control (debuffs)'): section=qt
                    elif qs=='List Paragraph': bullets.append({'section':section,'text':qt,'rank5':rank5(qt)})
                    elif not desc: desc=qt
                j+=1
            heads.append({'sourceClass':cls,'sourceDiscipline':disc or 'General','sourceName':strip_paren(t),'desc':desc,'effects':bullets})
    if len(heads)!=len(doc.tables): raise SystemExit(f'heading/table mismatch {len(heads)} != {len(doc.tables)}')
    for idx,(r,tbl) in enumerate(zip(heads,doc.tables),1):
        cells={}
        for row in tbl.rows:
            for c in row.cells:
                txt=' '.join(c.text.split())
                if ':' in txt:
                    k,v=txt.split(':',1); cells[k.strip()]=v.strip()
        r.update(sourceIndex=idx,type=cells.get('Tipo',''),castTime=lastnum(cells.get('T. lanzamiento',''),0),
                 gcd=cells.get('GCD','Corto') or 'Corto',cooldown=lastnum(cells.get('Cooldown',''),0),
                 durationText=cells.get('Duración','-'),duration=lastnum(cells.get('Duración','-'),0),
                 manaRank5=(lambda x: (sum(abs(float(n.replace(',','.'))) for n in NUM.findall(x))/len(NUM.findall(x))) if NUM.findall(x) else 0.0)((cells.get('Maná','') or '').split('/')[-1]),
                 range=lastnum(cells.get('Rango',''),0),area=lastnum(cells.get('Área',''),0),
                 weaponScale=boolish(cells.get('Escala con arma','')),blockable=boolish(cells.get('Bloqueable','')),resistible=boolish(cells.get('Resistible','')))
    return heads

def load_runtime(path:Path):
    out=[]
    rx=re.compile(r'^\s*add\((\{.*\})\);\s*$')
    for ln in path.read_text(encoding='utf8').splitlines():
        m=rx.match(ln)
        if m: out.append(json.loads(m.group(1)))
    return out

def walk_fx(x):
    if isinstance(x,list):
        for y in x: yield from walk_fx(y)
    elif isinstance(x,dict):
        if 'type' in x: yield x
        for k in ('effects','selfEffects','companionEffects','then','otherwise','onTrigger'):
            if k in x: yield from walk_fx(x[k])

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('source',type=Path); ap.add_argument('--runtime',type=Path,default=Path('js/data/powerLibrary.js')); ap.add_argument('--resolver',type=Path,default=Path('js/combat/resolver.js'))
    a=ap.parse_args(); src=parse_source(a.source); rt=load_runtime(a.runtime)
    failures=[]; gates=[]
    def gate(name,ok,detail=''):
        gates.append((name,ok,detail));
        if not ok: failures.append(name+(' — '+detail if detail else ''))

    valid=[r for r in src if r['sourceName'].strip() and not r['sourceName'].lower().startswith('undefined')]
    excluded=[r for r in src if r not in valid]
    gate('Fuente 320 / válidos 290 / placeholders 30',len(src)==320 and len(valid)==290 and len(excluded)==30,f'{len(src)}/{len(valid)}/{len(excluded)}')
    expected=[]
    for r in valid:
        for cid in TARGETS[r['sourceClass']]: expected.append((r['sourceIndex'],cid))
    actual=[(p['sourceIndex'],p['classId']) for p in rt]
    gate('Asignaciones exactas 410, sin extras ni faltantes',len(actual)==410 and sorted(actual)==sorted(expected),f'{len(actual)} runtime')
    counts={c:sum(1 for _,x in actual if x==c) for c in ['devastador','guardian','centinela','rastreador','arcanista','vinculador']}
    gate('Conteo por subclase 65/65/65/65/75/75',counts=={'devastador':65,'guardian':65,'centinela':65,'rastreador':65,'arcanista':75,'vinculador':75},str(counts))

    by_src={r['sourceIndex']:r for r in valid}; timing_bad=[]; metadata_bad=[]; name_bad=[]; effect_text_bad=[]; costs_bad=[]
    for p in rt:
        r=by_src[p['sourceIndex']]; m=p.get('sourceMechanics',{})
        if p['name']==r['sourceName'] or norm(r['sourceName']) not in norm(p['name']): name_bad.append(f"{p['id']}:{p['name']}←{r['sourceName']}")
        if not p['name'].endswith(SUFFIX[r['sourceClass']]): name_bad.append(f"{p['id']}:suffix")
        pairs=[('sourceName',r['sourceName']),('sourceClass',r['sourceClass']),('sourceDiscipline',r['sourceDiscipline']),('type',r['type']),('gcd',r['gcd']),('durationText',r['durationText']),('weaponScale',r['weaponScale']),('blockable',r['blockable']),('resistible',r['resistible'])]
        for k,v in pairs:
            if m.get(k)!=v: metadata_bad.append(f"{p['id']}:{k} {m.get(k)!r}!={v!r}")
        for k in ('castTime','cooldown','duration','manaRank5','range','area'):
            if abs(float(m.get(k,0))-float(r[k]))>1e-9: metadata_bad.append(f"{p['id']}:{k} {m.get(k)}!={r[k]}")
        expected_effects=r['effects']; got=m.get('effects',[])
        if got!=expected_effects: effect_text_bad.append(p['id'])
        if p.get('flags',{}).get('passive'):
            if p.get('cost')!=0: costs_bad.append(p['id']+':passive-cost')
        elif abs(float(p.get('cost',0))-float(r['manaRank5']))>1e-9: costs_bad.append(f"{p['id']}:{p.get('cost')}!={r['manaRank5']}")
        if abs(float(p.get('castTime',0))-float(r['castTime']))>1e-9 or abs(float(p.get('cooldown',0))-float(r['cooldown']))>1e-9: timing_bad.append(p['id'])
        expected_g=GCD.get(r['gcd'],'short')
        if p.get('gcd')!=expected_g: timing_bad.append(p['id']+':gcd')
    gate('Nombres reconocibles: fuente + variación sistemática',not name_bad,'; '.join(name_bad[:8]))
    gate('Metadatos fuente campo por campo intactos',not metadata_bad,'; '.join(metadata_bad[:8]))
    gate('Todos los bullets y rank-5 permanecen auditables',not effect_text_bad,', '.join(effect_text_bad[:8]))
    gate('Cast / GCD / cooldown runtime fieles',not timing_bad,', '.join(timing_bad[:8]))
    gate('Maná runtime = coste rank-5 fuente',not costs_bad,'; '.join(costs_bad[:8]))

    bad_models=[p['id'] for p in rt if p.get('damageModel')!='source-rank5-exact' or not p.get('sourceDerived')]
    gate('Sin poderes re-balanceados/fixed-pure en catálogo fuente',not bad_models,', '.join(bad_models[:8]))
    raw=a.runtime.read_text(encoding='utf8')
    gate('Sin placeholders undefined ni reglas opacas rule_*',('undefined (' not in raw.lower() and 'rule_' not in raw), 'rule_/undefined encontrado' if ('rule_' in raw or 'undefined (' in raw.lower()) else '')

    stubs=[]; ranges=[]; types=set()
    for p in rt:
        fx=list(walk_fx([p.get('effects',[]),p.get('selfEffects',[]),p.get('companionEffects',[])])); types.update(x.get('type') for x in fx if x.get('type'))
        if not p.get('flags',{}).get('passive') and not fx: stubs.append(p['id'])
        for e in fx:
            for lo,hi in (('min','max'),('pctMin','pctMax')):
                if lo in e and hi in e and float(e[lo])>float(e[hi]): ranges.append(f"{p['id']}:{e['type']} {e[lo]}>{e[hi]}")
    gate('Todo poder lanzable tiene efecto ejecutable',not stubs,', '.join(stubs[:8]))
    gate('Todos los rangos de daño/porcentaje están ordenados',not ranges,'; '.join(ranges[:8]))
    resolver=a.resolver.read_text(encoding='utf8'); missing=[]
    for t in sorted(types):
        if not re.search(r'\n\s*'+re.escape(t)+r'\s*:\s*function\s*\(',resolver): missing.append(t)
    gate('Cada type generado tiene handler en resolver',not missing,', '.join(missing))

    # Explicit iconic adversarial anchors.
    iconic={245:('Meteorito del Vacío',2,15,220,30),261:('Bola de fuego del Vacío',1.5,15,185,25),262:('Explosión de hielo del Vacío',1,20,185,25),263:('Relámpago del Vacío',2,20,185,25)}
    ib=[]
    for idx,vals in iconic.items():
        p=next((x for x in rt if x['classId']=='arcanista' and x['sourceIndex']==idx),None)
        if not p or (p['name'],p['castTime'],p['cooldown'],p['cost'],p['range'])!=vals: ib.append(str(idx))
    gate('Anclas Brujo: Meteorito/Fuego/Hielo/Relámpago exactas',not ib,', '.join(ib))

    # Special mechanics whose semantic meaning cannot be proven from a scalar field.
    def srcpow(idx,cid): return next(x for x in rt if x['sourceIndex']==idx and x['classId']==cid)
    special=[]
    s256=srcpow(256,'arcanista'); special.append(any(e.get('type')=='sourceManaDrainDot' and e.get('duration')==30 for e in s256['effects']))
    s153=srcpow(153,'rastreador'); special.append(any(e.get('type')=='companionProtectOwner' and abs(e.get('redirectPct',0)-.3)<1e-9 for e in s153['selfEffects']))
    s292=srcpow(292,'vinculador'); special.append(any(e.get('type')=='companionProtectOwner' and abs(e.get('redirectPct',0)-.5)<1e-9 for e in s292['selfEffects']))
    s159=srcpow(159,'rastreador'); special.append(any(e.get('type')=='companionAoE' and e.get('radius')==10 for e in s159['selfEffects']))
    gate('Mecánicas especiales: drains/pets/enlaces no degradadas',all(special),str(special))

    # Semantic edge cases found by adversarial review. These are intentionally
    # independent of the generator heuristics so a future parser regression is loud.
    semantic=[]
    s251=srcpow(251,'arcanista')
    semantic.append(s251.get('target')=='enemy' and abs(s251.get('range',0)-20)<1e-9 and any(e.get('type')=='sourceDrain' and e.get('min')==450 and e.get('max')==600 for e in s251.get('effects',[])))
    s281=srcpow(281,'vinculador')
    semantic.append(any(e.get('type')=='sourceHeal' and e.get('min')==700 and not e.get('percentOfMax') for e in s281.get('effects',[])) and any(e.get('type')=='sourceHeal' and e.get('min')==5 and e.get('percentOfMax') for e in s281.get('effects',[])))
    s287=srcpow(287,'vinculador')
    semantic.append(any(e.get('type')=='sourceHeal' and e.get('min')==60 and e.get('percentOfMax') for e in s287.get('effects',[])))
    s156=srcpow(156,'rastreador')
    semantic.append(any(e.get('type')=='companionRevive' and e.get('hpPct')==1.0 for e in s156.get('selfEffects',[])) and not any(e.get('type')=='tameCreature' for e in walk_fx([s156.get('effects',[]),s156.get('selfEffects',[])])))
    s17=srcpow(17,'devastador')
    semantic.append(any(e.get('type')=='purge' and e.get('count')==2 and e.get('chance')==1.0 for e in s17.get('effects',[])))
    gate('Bordes semánticos: Vampirismo/curas %/revivir mascota/purga exactos',all(semantic),str(semantic))

    opaque=[]
    for p in rt:
        for e in walk_fx([p.get('effects',[]),p.get('selfEffects',[]),p.get('companionEffects',[])]):
            for k in ((e.get('data') or {}).get('sourceStats') or {}):
                if k.startswith('source_'): opaque.append(f"{p['id']}:{k}")
    gate('Cero fallbacks semánticos opacos source_*',not opaque,', '.join(opaque[:8]))

    # Targeting/area semantics: a scalar-perfect power is still wrong if it
    # buffs the wrong side, centers the area incorrectly or self-casts an ally link.
    targeting=[]
    s57=srcpow(57,'guardian'); targeting.append(s57.get('target')=='aoeSelf' and s57.get('affects')=='alliesAndSelf' and s57.get('radius')==10)
    s59=srcpow(59,'guardian'); targeting.append(s59.get('target')=='aoeSelf' and s59.get('affects')=='alliesAndSelf' and s59.get('radius')==10)
    s60=srcpow(60,'guardian'); targeting.append(any(e.get('type')=='aura' and e.get('affects')=='alliesAndSelf' and e.get('radius')==6 and e.get('duration')==120 for e in s60.get('selfEffects',[])))
    s64=srcpow(64,'guardian'); targeting.append(s64.get('target')=='allyOrSelf' and any(e.get('effect')=='damageRedirect' and e.get('data',{}).get('redirectPct')==1.0 for e in s64.get('effects',[])))
    s219=srcpow(219,'vinculador'); targeting.append(any(e.get('type')=='aura' and e.get('affects')=='alliesAndSelf' and e.get('radius')==6 for e in s219.get('selfEffects',[])))
    s225=srcpow(225,'arcanista'); targeting.append(any(e.get('type')=='aura' and e.get('affects')=='enemies' and e.get('radius')==10 for e in s225.get('selfEffects',[])))
    s260=srcpow(260,'arcanista'); targeting.append(any(e.get('type')=='aura' and e.get('affects')=='enemies' and e.get('radius')==10 for e in s260.get('selfEffects',[])))
    s289=srcpow(289,'vinculador'); targeting.append(s289.get('target')=='targetArea' and s289.get('affects')=='allies' and s289.get('range')==20 and s289.get('radius')==6)
    s306=srcpow(306,'vinculador'); s307=srcpow(307,'vinculador'); targeting.append(s306.get('target')=='allyOrSelf' and s307.get('target')=='allyOrSelf')
    gate('Objetivo, bando y centro de área preservados en poderes sensibles',all(targeting),str(targeting))

    typed=[]
    typed.append(any(e.get('data',{}).get('physicalDamageTakenPct')==-0.5 and 'magicalDamageTakenPct' not in e.get('data',{}) for e in s57.get('effects',[])))
    typed.append(any(e.get('data',{}).get('magicalDamageTakenPct')==-0.5 and 'physicalDamageTakenPct' not in e.get('data',{}) for e in s59.get('effects',[])))
    typed.append(any(e.get('data',{}).get('magicalDamageTakenPct')==-0.3 for e in s306.get('effects',[])))
    typed.append(any(e.get('data',{}).get('physicalDamageTakenPct')==-0.3 for e in s307.get('effects',[])))
    gate('Resistencias físicas/mágicas conservan su canal, sin defensa genérica falsa',all(typed),str(typed))

    counters=[]
    s124=srcpow(124,'centinela'); counters.append(any(e.get('effect')=='sourceRetaliation' and e.get('data',{}).get('returnPct')==0.9 and e.get('duration')==15 for e in s124.get('effects',[])))
    s305=srcpow(305,'vinculador'); counters.append(any(e.get('effect')=='sourceDamageReflect' and e.get('data',{}).get('returnPct')==0.3 and e.get('duration')==30 for e in s305.get('effects',[])))
    gate('Represalia y Espejo del karma no se degradan al reflector mágico legacy',all(counters),str(counters))

    print('POWER PARITY ARBITER v0.13')
    for name,ok,detail in gates: print(('PASS' if ok else 'FAIL')+' | '+name+((' | '+detail) if detail else ''))
    print(f'\nGATES: {sum(1 for _,o,_ in gates if o)}/{len(gates)}')
    print('ARBITER: '+('APROBADO' if not failures else 'RECHAZADO'))
    if failures:
        print('\nFALLOS:'); [print(' - '+x) for x in failures]
        return 1
    return 0
if __name__=='__main__': sys.exit(main())
